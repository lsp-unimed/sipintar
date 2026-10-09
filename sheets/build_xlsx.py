"""Membangun Database-SIPINTAR-LSP-UNIMED.xlsx (template untuk diunggah ke Google Drive → Google Sheets).
Urutan & nama kolom HARUS sama dengan SHEETS di backend/Code.gs."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
import sys, json, os
DATA = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'data_lsp.json'), encoding='utf-8'))

OUT = sys.argv[1] if len(sys.argv) > 1 else 'Database-SIPINTAR-LSP-UNIMED.xlsx'
F = Font(name='Arial', size=10, color='000000')
FB = Font(name='Arial', size=10, bold=True, color='000000')
HEAD_FILL = PatternFill('solid', fgColor='D9DDE4')
ISI_FILL = PatternFill('solid', fgColor='FFF2A8')   # sel yang wajib disesuaikan
THIN = Side(style='thin', color='B7BDC7')
N = 1000  # baris yang diformat teks

SHEETS = {
  'Pengaturan': ['kunci', 'nilai', 'keterangan'],
  'Skema': ['id_skema', 'kode_skema', 'nama_skema', 'jenis_skema', 'jumlah_unit', 'persyaratan', 'biaya', 'link_dokumen', 'status'],
  'TUK': ['id_tuk', 'nama_tuk', 'jenis_tuk', 'alamat', 'penanggung_jawab', 'kontak', 'status'],
  'Asesor': ['id_asesor', 'nama_asesor', 'no_reg_met', 'skema', 'email', 'hp', 'status'],
  'Jadwal': ['id_jadwal', 'id_skema', 'tanggal', 'waktu', 'id_tuk', 'kuota', 'batas_daftar', 'status', 'keterangan'],
  'Pendaftaran': ['no_reg', 'waktu_daftar', 'id_jadwal', 'id_skema', 'nama', 'nik', 'nim', 'tempat_lahir', 'tanggal_lahir',
    'jenis_kelamin', 'email', 'hp', 'alamat', 'instansi', 'pendidikan', 'pekerjaan', 'tujuan_asesmen',
    'file_ktp', 'file_foto', 'file_ijazah', 'file_apl02', 'file_pendukung',
    'status_verifikasi', 'catatan_verifikasi', 'tgl_verifikasi',
    'status_jadwal', 'id_asesor', 'id_tuk', 'tanggal_asesmen', 'waktu_asesmen',
    'status_asesmen', 'catatan_asesmen',
    'rekomendasi', 'tgl_hasil', 'link_surat_hasil', 'catatan_hasil',
    'status_sertifikat', 'no_sertifikat', 'tgl_serah', 'penerima', 'diperbarui',
    'file_apl01', 'status_kelengkapan', 'catatan_kelengkapan', 'cek_berkas', 'tgl_kelengkapan',
    'verifikasi_langsung', 'cek_persyaratan', 'rekomendasi_apl01'],
  'Keluhan': ['no_tiket', 'waktu', 'nama', 'email', 'hp', 'no_reg', 'kategori', 'isi', 'status', 'tindak_lanjut', 'tgl_selesai', 'petugas'],
  'Layanan': ['no_layanan', 'waktu', 'jenis', 'nama', 'email', 'hp', 'no_reg', 'no_sertifikat', 'skema', 'keterangan', 'file', 'status', 'catatan_petugas', 'tgl_selesai', 'petugas'],
  'Survei': ['waktu', 'no_reg', 'skor_informasi', 'skor_administrasi', 'skor_asesmen', 'skor_petugas', 'skor_keseluruhan', 'saran'],
  'Dokumen': ['id_dok', 'nomor', 'judul', 'kategori', 'id_skema', 'link', 'status'],
  'Log': ['waktu', 'aktor', 'peran', 'langkah_sop', 'aksi', 'ref', 'detail'],
  'Akses': ['waktu', 'aktor', 'peran', 'langkah_sop', 'aksi', 'ref', 'detail'],
  'Pengguna': ['username', 'nama', 'peran', 'salt', 'password_hash', 'aktif'],
}
WIDTH = {'nama_skema': 42, 'persyaratan': 60, 'nilai': 70, 'keterangan': 48, 'judul': 52, 'alamat': 36, 'nama_tuk': 36,
         'nama_asesor': 30, 'link': 40, 'link_dokumen': 40, 'aksi': 50, 'detail': 40, 'isi': 50, 'nama': 26, 'kunci': 30}

wb = Workbook()
pet = wb.active
pet.title = 'Petunjuk'

def sheet(name):
    ws = wb.create_sheet(name)
    head = SHEETS[name]
    ws.append(head)
    for c in ws[1]:
        c.font = FB; c.fill = HEAD_FILL; c.border = Border(bottom=THIN)
        c.alignment = Alignment(vertical='center')
    ws.freeze_panes = 'A2'
    for i, h in enumerate(head, 1):
        col = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col].width = WIDTH.get(h, max(14, len(h) + 4))
        for r in range(2, N + 1):
            ws.cell(row=r, column=i).number_format = '@'
    return ws

NEXT = {}  # baris berikutnya per sheet (jangan pakai ws.append: 1000 baris sudah diformat '@')

def put(ws, rows, fill_cols=()):
    head = SHEETS[ws.title]
    for row in rows:
        r = NEXT.get(ws.title, 2); NEXT[ws.title] = r + 1
        for i, h in enumerate(head, 1):
            c = ws.cell(row=r, column=i, value=str(row.get(h, '')))
            c.font = F; c.number_format = '@'
            c.alignment = Alignment(wrap_text=h in ('persyaratan', 'nilai', 'keterangan'), vertical='top')
            if h in fill_cols: c.fill = ISI_FILL

def dv(ws, col_name, options, rng=N):
    head = SHEETS[ws.title]
    col = ws.cell(row=1, column=head.index(col_name) + 1).column_letter
    v = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=True)
    v.add(f'{col}2:{col}{rng}')
    ws.add_data_validation(v)

for n in SHEETS: sheet(n)

# ---------- Pengaturan ----------
put(wb['Pengaturan'], [
  dict(kunci='nama_lsp', nilai='LSP Universitas Negeri Medan', keterangan='Nama resmi LSP'),
  dict(kunci='nama_singkat', nilai='LSP UNIMED', keterangan='Dipakai di subjek email'),
  dict(kunci='tagline', nilai='Sistem Informasi Pemantauan dan Layanan Terintegrasi LSP Universitas Negeri Medan', keterangan=''),
  dict(kunci='deskripsi', nilai='Layanan digital sertifikasi kompetensi LSP Universitas Negeri Medan.', keterangan=''),
  dict(kunci='nomor_lisensi', nilai='', keterangan='Isi nomor lisensi BNSP'),
  dict(kunci='alamat', nilai='Jl. William Iskandar Ps. V, Medan Estate, Deli Serdang, Sumatera Utara 20221', keterangan='Periksa / sesuaikan'),
  dict(kunci='email', nilai='lspunimed@unimed.ac.id', keterangan='Email resmi LSP (juga alamat balasan email otomatis)'),
  dict(kunci='telepon', nilai='', keterangan=''),
  dict(kunci='whatsapp', nilai='', keterangan='Format 62812xxxx (tanpa + dan spasi)'),
  dict(kunci='jam_layanan', nilai='Senin–Jumat, 08.00–16.00 WIB', keterangan=''),
  dict(kunci='pengumuman', nilai='Silakan pelajari jadwal pelaksanaan uji kompetensi beserta persyaratan setiap skema sertifikasi sebelum melakukan pendaftaran. Informasi pada setiap tahapan pendaftaran akan disampaikan melalui email yang Saudara daftarkan, sehingga pastikan alamat email tersebut aktif. Perkembangan status permohonan juga dapat dipantau sewaktu-waktu melalui menu Lacak Permohonan di laman ini dengan menggunakan Nomor Registrasi dan email terdaftar.', keterangan='Tampil di beranda'),
  dict(kunci='link_template_apl02', nilai='', keterangan='Opsional. Template APL-02 per skema diisi di sheet Dokumen'),
  dict(kunci='kirim_email', nilai='YA', keterangan='YA = kirim email otomatis (bukti daftar, verifikasi, jadwal, hasil, sertifikat siap). TIDAK = matikan'),
  dict(kunci='link_contoh_apl01', nilai='', keterangan='Opsional. Contoh umum APL-01 bila contoh per skema di sheet Dokumen belum diisi'),
  dict(kunci='link_contoh_apl02', nilai='', keterangan='Opsional. Contoh umum APL-02 bila contoh per skema di sheet Dokumen belum diisi'),
  dict(kunci='info_verifikasi_langsung', nilai='Verifikasi langsung berkas asli di Sekretariat LSP UNIMED, Senin–Jumat pukul 08.00–15.00 WIB, paling lambat 15 Oktober 2026.', keterangan='Tempat & waktu verifikasi langsung (tampil ke asesi dan di email)'),
  dict(kunci='info_pengambilan_sertifikat', nilai='Sertifikat diambil di Sekretariat LSP pada jam layanan dengan membawa KTP asli. Pengambilan oleh orang lain wajib membawa surat kuasa.', keterangan=''),
], fill_cols=())
for r in (6, 9):
    wb['Pengaturan'].cell(row=r, column=2).fill = ISI_FILL

# ---------- Skema (16) dari dokumen skema bagian 9.1.2 ----------
SK = DATA['skema']
put(wb['Skema'], [dict(id_skema=f'SKM-{i:03d}', kode_skema='', nama_skema=k['nama'], jenis_skema=k['jenis'], jumlah_unit='',
    persyaratan='\n'.join(k['syarat']), biaya='', link_dokumen='', status='Aktif') for i, k in enumerate(SK, 1)], fill_cols=('kode_skema', 'jumlah_unit'))
dv(wb['Skema'], 'jenis_skema', ['KKNI', 'Okupasi', 'Klaster'])
dv(wb['Skema'], 'status', ['Aktif', 'Nonaktif'])
wb['Skema']['F2'].comment = Comment('Persyaratan diambil dari bagian 9.1.2 dokumen skema. Satu persyaratan per baris (Alt+Enter di Excel / Ctrl+Enter di Google Sheets).', 'SIPINTAR')
for i, k in enumerate(SK, 2):
    if k.get('catatan'):
        c = wb['Skema'].cell(row=i, column=6); c.fill = ISI_FILL; c.comment = Comment(k['catatan'], 'SIPINTAR')

# ---------- TUK (surat 071/LSP-UNIMED/X/2026) ----------
put(wb['TUK'], [dict(id_tuk=t[0], nama_tuk=t[1], jenis_tuk='Sewaktu', alamat='Universitas Negeri Medan', penanggung_jawab=t[2], kontak='', status='Aktif') for t in DATA['tuk']])
dv(wb['TUK'], 'jenis_tuk', ['Sewaktu', 'Tempat Kerja', 'Mandiri'])
dv(wb['TUK'], 'status', ['Aktif', 'Nonaktif'])

# ---------- Asesor (surat 071/LSP-UNIMED/X/2026) ----------
put(wb['Asesor'], [dict(id_asesor=f'ASR-{i:03d}', nama_asesor=k['asesor'][0], no_reg_met='MET.' + k['asesor'][1], skema=f'SKM-{i:03d}', email='', hp='', status='Aktif') for i, k in enumerate(SK, 1)])
dv(wb['Asesor'], 'status', ['Aktif', 'Nonaktif'])
wb['Asesor']['D2'].comment = Comment('ID skema yang diampu, pisahkan koma. Contoh: SKM-001,SKM-004', 'SIPINTAR')

# ---------- Jadwal: 16 skema, 17 Oktober 2026 ----------
put(wb['Jadwal'], [dict(id_jadwal=f'JDW-{i:03d}', id_skema=f'SKM-{i:03d}', tanggal='2026-10-17', waktu='08.00 WIB – selesai', id_tuk=k['tuk'],
    kuota='0', batas_daftar='2026-10-14', status='Dibuka', keterangan='Uji kompetensi perdana (penyaksian BNSP)') for i, k in enumerate(SK, 1)], fill_cols=('batas_daftar',))
dv(wb['Jadwal'], 'status', ['Dibuka', 'Ditutup', 'Selesai'])
wb['Jadwal']['F2'].comment = Comment('0 = tanpa batas. Kuota tidak ditampilkan di situs publik.', 'SIPINTAR')
wb['Jadwal']['G2'].comment = Comment('Batas pendaftaran (format yyyy-mm-dd). Sesuaikan.', 'SIPINTAR')

# ---------- Dokumen: umum + per skema ----------
dok = [
  ('XVII/SOP-PKS', 'SOP Pelayanan Kegiatan Sertifikasi', 'SOP', ''),
  ('PBNSP 201', 'Persyaratan Umum Lembaga Sertifikasi Profesi', 'Acuan', ''),
  ('PBNSP 202', 'Pelaksanaan Sertifikasi Kompetensi', 'Acuan', ''),
  ('PBNSP 206', 'Sistem Manajemen Mutu LSP', 'Acuan', ''),
  ('ISO/IEC 17024:2012', 'Conformity Assessment – General Requirements for Bodies Operating Certification of Persons', 'Acuan', ''),
]
for i, k in enumerate(SK, 1):
    dok.append(('FR.APL.01', 'Permohonan Sertifikasi Kompetensi (APL-01) — ' + k['nama'], 'Formulir', f'SKM-{i:03d}'))
    dok.append(('FR.APL.02', 'Asesmen Mandiri (APL-02) — ' + k['nama'], 'Formulir', f'SKM-{i:03d}'))
    dok.append(('FR.APL.01', 'Contoh pengisian APL-01 — ' + k['nama'], 'Contoh Pengisian', f'SKM-{i:03d}'))
    dok.append(('FR.APL.02', 'Contoh pengisian APL-02 — ' + k['nama'], 'Contoh Pengisian', f'SKM-{i:03d}'))
    dok.append(('FR.SKEMA-02', 'Skema Sertifikasi ' + k['nama'], 'Skema', f'SKM-{i:03d}'))
put(wb['Dokumen'], [dict(id_dok=f'DOK-{k:03d}', nomor=a, judul=b, kategori=c, id_skema=d, link='', status='Aktif') for k, (a, b, c, d) in enumerate(dok, 1)], fill_cols=('link',))
dv(wb['Dokumen'], 'kategori', ['SOP', 'Acuan', 'Formulir', 'Contoh Pengisian', 'Skema', 'Panduan', 'Lainnya'])
dv(wb['Dokumen'], 'status', ['Aktif', 'Nonaktif'])
wb['Dokumen']['E2'].comment = Comment('Kosong = dokumen umum (semua skema). Isi ID skema (mis. SKM-003) untuk formulir khusus skema.', 'SIPINTAR')
wb['Dokumen']['F2'].comment = Comment('Link Google Drive dengan akses "Siapa saja yang memiliki link".', 'SIPINTAR')

# ---------- Validasi sheet transaksi ----------
P = wb['Pendaftaran']
dv(P, 'status_kelengkapan', ['Menunggu Pemeriksaan', 'Lengkap', 'Belum Lengkap'])
dv(P, 'verifikasi_langsung', ['YA', 'TIDAK'])
dv(P, 'rekomendasi_apl01', ['Diterima', 'Tidak diterima'])
dv(P, 'status_verifikasi', ['Menunggu Verifikasi', 'Memenuhi Syarat', 'Perlu Perbaikan', 'Tidak Memenuhi Syarat'])
dv(P, 'status_jadwal', ['Belum Dijadwalkan', 'Terjadwal'])
dv(P, 'status_asesmen', ['Belum', 'Dokumen Siap', 'Hadir', 'Tidak Hadir', 'Selesai'])
dv(P, 'rekomendasi', ['Kompeten', 'Belum Kompeten'])
dv(P, 'status_sertifikat', ['Belum Terbit', 'Diajukan ke BNSP', 'Siap Diambil', 'Sudah Diserahkan'])
for n in ('Keluhan', 'Layanan'): dv(wb[n], 'status', ['Diterima', 'Diproses', 'Selesai', 'Ditolak'])
dv(wb['Pengguna'], 'peran', ['Admin', 'Sekretariat LSP', 'Bagian Administrasi', 'Bagian Sertifikasi', 'Bagian Manajemen Mutu'])
dv(wb['Pengguna'], 'aktif', ['YA', 'TIDAK'])

# ---------- Petunjuk ----------
lines = [
  ('Database SIPINTAR LSP UNIMED', True),
  ('Basis data layanan sertifikasi untuk penyaksian uji kompetensi BNSP. Dibaca dan ditulis oleh Apps Script (Code.gs).', False),
  ('', False),
  ('Cara memasang', True),
  ('1. Unggah file ini ke Google Drive, klik kanan → Buka dengan → Google Spreadsheet. Lalu File → Simpan sebagai Google Spreadsheet.', False),
  ('2. Di spreadsheet hasil konversi: Ekstensi → Apps Script. Tempel isi Code.gs, simpan.', False),
  ('3. Jalankan fungsi setup() sekali dan beri izin. Password akun "admin" muncul di Log eksekusi.', False),
  ('4. Terapkan → Deployment baru → Aplikasi web (Jalankan sebagai: Saya, Akses: Siapa saja). Salin URL /exec ke assets/js/config.js.', False),
  ('', False),
  ('Yang wajib diisi sebelum dipakai (sel berwarna kuning)', True),
  ('Skema: 16 skema dan persyaratannya sudah diisi dari bagian 9.1.2 dokumen skema. Lengkapi kode skema dan jumlah unit; periksa sel persyaratan berwarna kuning (ada catatan).', False),
  ('Jadwal: 16 jadwal sudah disiapkan untuk 17 Oktober 2026 dengan TUK sesuai surat 071/LSP-UNIMED/X/2026. Periksa batas_daftar.', False),
  ('TUK dan Asesor: sudah diisi dari surat 071/LSP-UNIMED/X/2026. Lengkapi email/HP bila perlu.', False),
  ('Dokumen: link Google Drive untuk SOP, APL-01, APL-02, dan contoh pengisian APL-01/APL-02 per skema (akses "Siapa saja yang memiliki link").', False),
  ('Pengaturan: nomor lisensi, email resmi, telepon.', False),
  ('', False),
  ('Aturan penting', True),
  ('Jangan mengubah nama sheet, nama kolom di baris 1, atau urutan kolom.', False),
  ('Tanggal ditulis teks yyyy-mm-dd, contoh 2026-10-17. Kolom diformat teks agar NIK dan nomor HP tidak berubah.', False),
  ('Sheet Pendaftaran, Keluhan, Layanan, Survei, Log, dan Pengguna diisi otomatis oleh sistem. Pengguna dibuat oleh setup() dan panel admin.', False),
  ('Bagikan spreadsheet ini hanya kepada admin LSP. Admin lain bekerja lewat panel web sesuai perannya.', False),
]
for t, b in lines:
    pet.append([t])
    pet.cell(row=pet.max_row, column=1).font = Font(name='Arial', size=12 if b and pet.max_row == 1 else 10, bold=b, color='000000')
pet.column_dimensions['A'].width = 120

for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 100
wb.save(OUT)
print('OK', OUT)
